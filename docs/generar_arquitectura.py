"""
Genera el documento Word de arquitectura para la Tarea 9 de SENATI.
"""
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

doc = Document()

# ── Estilos base ──
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

def set_cell_shading(cell, color):
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(qn('w:shd'), {
        qn('w:val'): 'clear',
        qn('w:color'): 'auto',
        qn('w:fill'): color,
    })
    shading.append(shading_elem)

def add_heading_styled(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x03, 0x69, 0xA1)
    return h

def add_card(text, title=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    if title:
        run = p.add_run(title + ": ")
        run.bold = True
        run.font.color.rgb = RGBColor(0x03, 0x69, 0xA1)
        run.font.size = Pt(12)
    run2 = p.add_run(text)
    run2.font.size = Pt(11)
    return p

def add_blue_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(10)
        set_cell_shading(cell, '0284C7')
    # Data
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
            if r % 2 == 0:
                set_cell_shading(cell, 'F0F9FF')
    doc.add_paragraph()
    return table


# ═══════════════════════════════════════════════════════════════
# PORTADA
# ═══════════════════════════════════════════════════════════════
for _ in range(4):
    doc.add_paragraph()

title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_p.add_run('ARQUITECTURA DEL SISTEMA')
run.bold = True
run.font.size = Pt(28)
run.font.color.rgb = RGBColor(0x03, 0x69, 0xA1)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = sub_p.add_run('Sistema de Tienda Online - Ecommerce MVP')
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

doc.add_paragraph()

info_items = [
    ('Curso', 'Desarrollo de Sistemas Web'),
    ('Tarea', 'N° 9 - Documentación de Arquitectura'),
    ('Alumno', 'Jorge Dante Lucano Tello'),
    ('Profesor', 'Iván Calle'),
    ('Institución', 'SENATI'),
    ('Fecha', 'Junio 2026'),
]
for label, value in info_items:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'{label}: ')
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x03, 0x69, 0xA1)
    run2 = p.add_run(value)
    run2.font.size = Pt(13)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# 1. INTRODUCCIÓN
# ═══════════════════════════════════════════════════════════════
add_heading_styled('1. Introducción', 1)

doc.add_paragraph(
    'Este documento describe la arquitectura técnica del sistema de tienda online desarrollado '
    'como proyecto de la Tarea 9. El sistema permite la venta de productos mediante un catálogo '
    'público, carrito de compras como invitado, checkout transaccional y un panel administrativo '
    'para la gestión de productos, stock y pedidos.'
)
doc.add_paragraph(
    'La arquitectura sigue el patrón monolito modular con separación clara entre frontend y backend, '
    'priorizando simplicidad, rápida salida al mercado y facilidad de mantenimiento.'
)

# ═══════════════════════════════════════════════════════════════
# 2. STACK TECNOLÓGICO
# ═══════════════════════════════════════════════════════════════
add_heading_styled('2. Stack Tecnológico', 1)

add_blue_table(
    ['Capa', 'Tecnología', 'Justificación'],
    [
        ['Frontend', 'Next.js 16 + React + Tailwind CSS', 'Renderizado híbrido (SSR/CSR), SEO y excelente DX'],
        ['Backend', 'FastAPI (Python 3.13)', 'Validaciones transaccionales, Swagger automático'],
        ['Base de Datos', 'SQLite (via SQLAlchemy)', 'Sin servidor externo, ideal para MVP y desarrollo'],
        ['ORM', 'SQLAlchemy 2.0 (async)', 'Mapeo objeto-relacional con soporte asíncrono'],
        ['Autenticación', 'JWT + bcrypt', 'Tokens sin estado, hashes seguros'],
        ['Validación', 'Pydantic v2', 'Schemas validados en tiempo de ejecución'],
        ['Estilos', 'Tailwind CSS', 'Diseño responsive sin escribir CSS manual'],
    ]
)

# ═══════════════════════════════════════════════════════════════
# 3. ARQUITECTURA
# ═══════════════════════════════════════════════════════════════
add_heading_styled('3. Diagrama de Arquitectura', 1)

doc.add_paragraph(
    'El sistema se compone de tres capas principales: Frontend (Next.js en puerto 3000), '
    'Backend (FastAPI en puerto 8000) y Base de Datos (SQLite).'
)

arch_text = (
    'NAVEGADOR (localhost:3000)\n'
    '    │\n'
    '    ├── Next.js 16 (Frontend)\n'
    '    │   ▸ Catálogo, Carrito, Checkout, Panel Admin\n'
    '    │   ▸ AuthGuard, CartContext, VariantSelector\n'
    '    │\n'
    '    └──▶ FastAPI (Backend :8000)\n'
    '         ▸ /api/v1/products, /checkout, /admin/*\n'
    '         ▸ JWT Auth, Validación de stock y precios\n'
    '              │\n'
    '              └──▶ SQLite (tienda.db)\n'
    '                   ▸ 14 tablas: products, orders, payments, etc.'
)
p = doc.add_paragraph()
run = p.add_run(arch_text)
run.font.name = 'Consolas'
run.font.size = Pt(9)
run.font.color.rgb = RGBColor(0x0C, 0x4A, 0x6E)

# ═══════════════════════════════════════════════════════════════
# 4. ZONAS
# ═══════════════════════════════════════════════════════════════
add_heading_styled('4. Zonas del Sistema', 1)

add_heading_styled('4.1 Zona Pública', 2)
doc.add_paragraph('Catálogo de productos activos, detalle con variantes (talla, color), carrito de compras (localStorage), checkout como invitado y confirmación de pedido. Sin autenticación requerida.')

add_heading_styled('4.2 Zona Administrativa', 2)
doc.add_paragraph('Login con JWT, dashboard con KPIs, CRUD de productos y variantes, gestión de imágenes, ajuste manual de stock, listado y detalle de pedidos, cambio de estado y confirmación de pagos. Requiere rol admin.')

# ═══════════════════════════════════════════════════════════════
# 5. BASE DE DATOS
# ═══════════════════════════════════════════════════════════════
add_heading_styled('5. Modelo de Base de Datos', 1)
doc.add_paragraph('El sistema utiliza 14 tablas organizadas en 3 módulos principales.')

add_heading_styled('5.1 Catálogo', 2)
add_blue_table(
    ['Tabla', 'Descripción'],
    [
        ['categories', 'Categorías de productos (slug único, soft delete)'],
        ['products', 'Producto padre: nombre, slug, descripción, material, marca'],
        ['product_variants', 'Unidad vendible: talla, color, SKU, precio, stock'],
        ['product_images', 'Imagen principal y galería por producto'],
    ]
)

add_heading_styled('5.2 Clientes y Acceso', 2)
add_blue_table(
    ['Tabla', 'Descripción'],
    [
        ['profiles', 'Usuarios del sistema con email y contraseña'],
        ['roles', 'Roles: admin, customer'],
        ['user_roles', 'Asignación perfil - rol'],
        ['customer_addresses', 'Direcciones guardadas'],
    ]
)

add_heading_styled('5.3 Operación Comercial', 2)
add_blue_table(
    ['Tabla', 'Descripción'],
    [
        ['orders', 'Cabecera del pedido con datos del cliente, totales y estado'],
        ['order_items', 'Líneas del pedido (snapshot de precio y variante)'],
        ['payments', 'Registro de pago (manual o simulado)'],
        ['payment_events', 'Historial de eventos de pago'],
        ['stock_adjustments', 'Bitácora de ajustes de stock'],
        ['order_status_history', 'Historial de transiciones de estado'],
    ]
)

add_heading_styled('5.4 Máquina de Estados del Pedido', 2)
doc.add_paragraph(
    'pendiente_pago → pagado → en_preparacion → listo_para_entrega → entregado\n'
    'pendiente_pago → cancelado\n'
    'pagado → cancelado / reembolsado'
)

# ═══════════════════════════════════════════════════════════════
# 6. API
# ═══════════════════════════════════════════════════════════════
add_heading_styled('6. API REST - Endpoints', 1)
doc.add_paragraph('Base URL: http://localhost:8000/api/v1 | Formato: JSON | Auth: JWT Bearer Token')

add_heading_styled('6.1 Endpoints Públicos', 2)
add_blue_table(
    ['Método', 'Ruta', 'Descripción'],
    [
        ['GET', '/health', 'Verificación de disponibilidad'],
        ['GET', '/products', 'Listado paginado con búsqueda y filtros'],
        ['GET', '/products/{slug}', 'Detalle completo con variantes e imágenes'],
        ['GET', '/categories', 'Lista de categorías activas'],
        ['GET', '/config', 'Configuración pública (pago, envío)'],
        ['POST', '/checkout/orders', 'Crear pedido como invitado'],
    ]
)

add_heading_styled('6.2 Endpoints Admin (JWT requerido)', 2)
add_blue_table(
    ['Método', 'Ruta', 'Descripción'],
    [
        ['POST', '/admin/auth/login', 'Iniciar sesión admin'],
        ['GET/POST/PUT/DEL', '/admin/products', 'CRUD de productos'],
        ['POST', '/admin/products/{id}/variants', 'Crear variante'],
        ['POST', '/admin/products/{id}/images', 'Subir imagen'],
        ['GET', '/admin/orders', 'Listar pedidos'],
        ['GET', '/admin/orders/{id}', 'Detalle de pedido'],
        ['POST', '/admin/orders/{id}/status-transitions', 'Cambiar estado'],
        ['POST', '/admin/orders/{id}/payments/...', 'Confirmar pago'],
        ['POST', '/admin/stock/adjustments', 'Ajustar stock'],
    ]
)

# ═══════════════════════════════════════════════════════════════
# 7. FRONTEND
# ═══════════════════════════════════════════════════════════════
add_heading_styled('7. Frontend - Componentes', 1)

add_heading_styled('7.1 Páginas Principales', 2)
add_blue_table(
    ['Ruta', 'Componente', 'Descripción'],
    [
        ['/', 'HomePage', 'Hero, categorías, productos destacados'],
        ['/catalogo', 'CatalogPage', 'Búsqueda, filtro, grid responsive'],
        ['/producto/[slug]', 'ProductDetail', 'Galería, variantes, CTA carrito'],
        ['/carrito', 'CartPage', 'Lista items, cantidades, resumen'],
        ['/checkout', 'CheckoutPage', 'Formulario, envío/recojo, pago'],
        ['/confirmacion', 'Confirmation', 'Nº de pedido, instrucciones'],
        ['/admin/login', 'AdminLogin', 'Formulario email + contraseña'],
        ['/admin', 'Dashboard', 'KPIs: productos, pedidos, stock bajo'],
        ['/admin/productos', 'AdminProducts', 'Tabla CRUD con formulario'],
        ['/admin/pedidos', 'AdminOrders', 'Listado, filtro, detalle'],
    ]
)

add_heading_styled('7.2 Componentes Reutilizables', 2)
doc.add_paragraph('• AuthGuard: Protege rutas /admin/*, redirige a login sin sesión')
doc.add_paragraph('• CartContext: Estado global del carrito con persistencia en localStorage')
doc.add_paragraph('• ProductCard: Tarjeta con imagen, precio formateado y badge de stock')
doc.add_paragraph('• VariantSelector: Botones de talla/color con filtro cruzado')
doc.add_paragraph('• ImageGallery: Imagen principal + thumbnails clickeables')
doc.add_paragraph('• ProductForm: Modal CRUD con gestión de variantes e imágenes')

# ═══════════════════════════════════════════════════════════════
# 8. SEGURIDAD
# ═══════════════════════════════════════════════════════════════
add_heading_styled('8. Seguridad', 1)

add_heading_styled('8.1 Autenticación', 2)
doc.add_paragraph('• JWT con algoritmo HS256, expira a las 8 horas')
doc.add_paragraph('• Contraseñas hasheadas con bcrypt')
doc.add_paragraph('• Middleware get_current_admin en todas las rutas protegidas')

add_heading_styled('8.2 Validaciones', 2)
doc.add_paragraph('• Precios recalculados en backend (nunca se confía en el frontend)')
doc.add_paragraph('• Stock validado atómicamente, no se permite stock negativo')
doc.add_paragraph('• Payloads validados con Pydantic v2')
doc.add_paragraph('• CORS restringido a localhost:3000')
doc.add_paragraph('• Errores sin stack traces en producción')

# ═══════════════════════════════════════════════════════════════
# 9. FLUJO DE COMPRA
# ═══════════════════════════════════════════════════════════════
add_heading_styled('9. Flujo Principal de Compra', 1)
doc.add_paragraph(
    '1. Catálogo → 2. Detalle Producto → 3. Seleccionar Variante (talla/color) → '
    '4. Agregar al Carrito → 5. Checkout (datos de contacto, envío, pago) → '
    '6. POST /checkout/orders → 7. Backend valida stock y precio → 8. Confirmación con nº de pedido'
)

# ═══════════════════════════════════════════════════════════════
# 10. CONCLUSIÓN
# ═══════════════════════════════════════════════════════════════
add_heading_styled('10. Conclusión', 1)

doc.add_paragraph(
    'La arquitectura implementada cumple con los requisitos de un MVP funcional para '
    'comercio electrónico, con separación clara entre zona pública y administrativa, '
    'backend con validaciones transaccionales, frontend mobile-first y persistencia '
    'en SQLite sin dependencia de servidores externos.'
)
doc.add_paragraph('Tecnologías utilizadas: Next.js 16, FastAPI, SQLite, Tailwind CSS, SQLAlchemy, JWT, bcrypt.')
doc.add_paragraph('El sistema está preparado para evolucionar a PostgreSQL y pasarela de pago real en futuras versiones.')

# ── FOOTER ──
doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('Jorge Dante Lucano Tello — Tarea 9 — SENATI — Junio 2026')
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run2 = p2.add_run('Profesor: Iván Calle | Curso: Desarrollo de Sistemas Web')
run2.font.size = Pt(10)
run2.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

# ── Guardar ──
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'ARQUITECTURA_TAREA9.docx')
doc.save(output_path)
print(f'Documento guardado en: {output_path}')
